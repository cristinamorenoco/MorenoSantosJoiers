import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Textos Web"

# --- Colores ---
COLOR_HEADER_BG = "1A1520"
COLOR_HEADER_FONT = "FFFFFF"
COLOR_PAGE_BG = "EDE0F5"
COLOR_PAGE_FONT = "1A1520"
COLOR_ROW_ODD = "FAF8FC"
COLOR_ROW_EVEN = "FFFFFF"
COLOR_CURRENT = "FFF3CD"
COLOR_PROPOSED = "D4EDDA"
COLOR_NOTES = "F8D7DA"

# --- Headers ---
headers = ["Página", "Sección", "Elemento", "Texto actual (web GitHub)", "Texto propuesto (Moreno Santos Joiers)", "Notas / Justificación"]
ws.append(headers)

header_row = ws[1]
for cell in header_row:
    cell.font = Font(bold=True, color=COLOR_HEADER_FONT, size=11)
    cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# --- Datos ---
data = [
    # ===================== HOME PAGE =====================
    ("Inicio", "HEAD / SEO", "Título de pestaña", "Moreno Santos — Joyería de Autor · Barcelona", "Moreno Santos Joiers — Joyería artesanal · Barcelona · Desde 1993", "Corregir año fundación (1993, no 2008). Añadir 'Joiers' (nombre real en catalán)."),
    ("Inicio", "Hero", "Tagline", "Barcelona · Joyería de Autor · Desde 2008", "Barcelona · Joiers artesans · Des de 1993", "La empresa se fundó en 1993. Incluir referencia bilingüe es natural en Barcelona."),
    ("Inicio", "Hero", "Descripción principal", "Cada pieza es una conversación entre el oro, la piedra y las manos del artesano. Creamos joyas que trascienden el tiempo.", "Cada joya nace de un diálogo entre el oro de 18 quilates, la piedra preciosa y más de treinta años de oficio artesanal. Creamos piezas que duran toda la vida.", "Destacar los 30 años de experiencia y el oro 18k como diferenciador real."),
    ("Inicio", "Hero", "CTA principal", "Explorar colección", "Descubre nuestras colecciones", "Más natural en español peninsular."),
    ("Inicio", "Hero", "CTA secundario", "Encargo a medida →", "Crea tu joya a medida →", "Más descriptivo y accionable."),
    ("Inicio", "Marquee", "Badge 1", "Oro 18 Quilates", "Oro 18 Quilates", "Correcto, mantener."),
    ("Inicio", "Marquee", "Badge 2", "Diamantes Certificados", "Diamantes y Gemas Naturales", "Ampliar para incluir todas las piedras preciosas naturales que trabajáis."),
    ("Inicio", "Marquee", "Badge 3", "Artesanía Barcelona", "Taller Propio · Barcelona", "Destacar que tenéis taller propio, es un diferenciador clave."),
    ("Inicio", "Marquee", "Badge 4", "Encargo Personal", "Alianzas de Boda", "Añadir uno de los servicios estrella."),
    ("Inicio", "Marquee", "Badge 5", "Gemas Naturales", "Reparación y Restauración", "Servicio real y diferenciador que ofrece la empresa."),
    ("Inicio", "Marquee", "Badge 6", "Edición Limitada", "Empresa Familiar · Desde 1993", "Posicionamiento clave: empresa familiar con más de 30 años."),
    ("Inicio", "Colecciones", "Título sección", "Nuestras colecciones — Piezas que cuentan historias", "Nuestras colecciones — Joyas que acompañan toda una vida", "Más emocional y vinculado a la filosofía de empresa duradera."),
    ("Inicio", "Colecciones", "Tarjeta 1", "Colección Signature / Diamantes", "Oro & Diamantes", "Simplificar y ser directo con los materiales reales."),
    ("Inicio", "Colecciones", "Tarjeta 2", "Alianzas", "Alianzas de Boda", "Especificar que son alianzas de boda, servicio estrella."),
    ("Inicio", "Colecciones", "Tarjeta 3", "Esmeraldas", "Piedras Preciosas", "Ampliar: trabajáis diamantes y piedras preciosas naturales en general."),
    ("Inicio", "Colecciones", "Tarjeta 4", "Edición limitada", "Plata & Relojes", "Incluir los productos de plata y relojes que también vendéis."),
    ("Inicio", "Historia", "Label sección", "Nuestra historia", "Nuestra historia — Más de 30 años", "Reforzar la trayectoria temporal real."),
    ("Inicio", "Historia", "Título", "Tradición artesana en el corazón de Barcelona", "Tradición artesana y familiar desde 1993", "La clave es empresa familiar + año real de fundación."),
    ("Inicio", "Historia", "Párrafo 1", "Moreno Santos nació de la pasión por la joyería de autor y el respeto por los materiales nobles. En nuestro atelier del Barrio Gótico, cada pieza se diseña y elabora a mano, uniendo técnicas centenarias con una visión contemporánea.", "Moreno Santos Joiers nació en 1993 como un proyecto familiar con una sola convicción: crear joyas únicas con oro de 18 quilates, diamantes y piedras preciosas naturales. En nuestro taller de Barcelona, cada pieza se diseña y elabora a mano, combinando técnica artesanal y sensibilidad contemporánea.", "Corregir año, añadir 'empresa familiar', mencionar materiales reales, eliminar 'Barrio Gótico' (dirección desconocida)."),
    ("Inicio", "Historia", "Párrafo 2", "Trabajamos exclusivamente con oro de 18 quilates, platino y gemas certificadas de origen ético. Porque la verdadera belleza no debería tener ningún coste oculto.", "Trabajamos con oro de 18 quilates, diamantes, piedras preciosas naturales, plata y relojes de las mejores marcas. Más de treinta años avalan nuestra pasión por el detalle y la calidad.", "Eliminar 'platino' si no es real. Añadir plata y relojes que sí vendéis. Mencionar las marcas."),
    ("Inicio", "Estadísticas", "Badge años", "15+ Años", "30+ Años", "Fundada en 1993: más de 30 años de trayectoria."),
    ("Inicio", "Estadísticas", "Stat 1", "850+ | Piezas creadas", "[Número real] | Piezas creadas", "Poner la cifra real de vuestra empresa o eliminarla si no disponéis del dato."),
    ("Inicio", "Estadísticas", "Stat 2", "38 | Países", "Empresa familiar | Desde 1993", "Sustituir dato posiblemente ficticio por mensaje más honesto y diferenciador."),
    ("Inicio", "Estadísticas", "Stat 3", "100% | Artesanía local", "100% | Taller propio", "Mantener el 100% artesanal pero especificar que tenéis taller propio."),
    ("Inicio", "Craft / Pilares", "Pilar 1", "Oro Certificado — 18 y 24 quilates — origen ético garantizado", "Oro de 18 Quilates — Diseños propios en oro 18k con acabado artesanal", "Ajustar a la realidad: solo 18k, y diseños propios es un diferenciador clave."),
    ("Inicio", "Craft / Pilares", "Pilar 2", "Gemas Naturales — Diamantes, esmeraldas y zafiros certificados", "Gemas Naturales — Diamantes y piedras preciosas naturales seleccionadas", "Correcto pero generalizar para incluir todas las piedras que trabajáis."),
    ("Inicio", "Craft / Pilares", "Pilar 3", "Hecho a Mano — Cada pieza elaborada en nuestro atelier", "Taller Propio — Cada joya creada y reparada en nuestro taller de Barcelona", "Incluir también la faceta de reparación, diferenciador real."),
    ("Inicio", "Craft / Pilares", "Pilar 4", "Encargo Personal — Diseñamos la joya de tu vida", "A tu Medida — Diseñamos, creamos y grabamos la joya de tu vida", "Incluir el servicio de grabación, que también ofrecéis."),
    ("Inicio", "Testimonio", "Cita", "Cada mañana me pongo mi alianza de Moreno Santos y siento que llevo un pedacito de arte conmigo. — Sofía R., clienta desde 2019 · Madrid", "[Testimonio real de cliente] — [Nombre], cliente desde [año]", "Sustituir por un testimonio real de cliente, mucho más creíble y auténtico."),
    ("Inicio", "CTA final", "Título", "Únete a nuestra comunidad", "Descubre todo lo que podemos crear juntos", "Más enfocado en la propuesta de valor artesanal y relación con el cliente."),
    ("Inicio", "CTA final", "Subtítulo", "Recibe lo extraordinario", "Una joya para cada momento", "Más cercano y menos pretencioso, adaptado a empresa familiar."),
    ("Inicio", "Footer", "Copyright", "© 2025 Moreno Santos · Barcelona", "© 2025 Moreno Santos Joiers · Barcelona · Desde 1993", "Añadir nombre completo y año fundación."),

    # ===================== NOSOTROS =====================
    ("Nosotros", "Hero", "Año fundación", "Fundada en Barcelona · 2008", "Empresa familiar fundada en Barcelona · 1993", "Corrección fundamental: año real y empresa familiar."),
    ("Nosotros", "Historia", "Título", "Nuestra historia — Más que joyas", "Más de treinta años haciendo joyas con alma", "Más directo, emocional y veraz."),
    ("Nosotros", "Historia", "Filosofía", "Moreno Santos nació de la convicción de que la joyería de verdad no se compra: se elige. Cada pieza que creamos lleva consigo una historia, un amor y una intención que trasciende el tiempo.", "Moreno Santos Joiers nació en 1993 de la convicción de que la joyería de verdad no se compra: se crea. Como empresa familiar, cada pieza que elaboramos lleva consigo una historia, un vínculo y una intención que trasciende el tiempo.", "Añadir 'empresa familiar' y 'que se crea' (taller propio). Corregir año."),
    ("Nosotros", "Manifiesto", "Cita", "Creamos para quienes entienden que el lujo verdadero no es el precio — es la singularidad, la artesanía y el vínculo emocional que une a una persona con su joya para siempre.", "Creamos para quienes entienden que una joya de verdad no es solo un objeto — es un recuerdo, una promesa y un legado que perdura toda la vida.", "Tono más humano y menos 'de lujo', más coherente con empresa familiar cercana."),
    ("Nosotros", "Valores", "Valor 1 título", "Artesanía sin concesiones", "Taller propio, sin intermediarios", "Destacar el taller propio como diferenciador real y concreto."),
    ("Nosotros", "Valores", "Valor 1 texto", "Nunca producimos en serie. Cada pieza pasa por las manos del orfebre desde el primer trazo hasta el último pulido. Eso nos hace más lentos — y mucho más buenos.", "Nunca producimos en serie. Contamos con taller propio en Barcelona donde cada joya nace desde el primer boceto hasta el acabado final. Eso nos hace más lentos — y mucho más cuidadosos.", "Sustituir 'orfebre' genérico por referencia al taller propio."),
    ("Nosotros", "Valores", "Valor 2 título", "Materiales de origen ético", "Materiales de calidad contrastada", "Más honesto si no tenéis certificación específica activa."),
    ("Nosotros", "Valores", "Valor 2 texto", "Solo trabajamos con proveedores certificados por el Proceso de Kimberley y estándares de minería responsable. La belleza no puede construirse sobre daño.", "Solo trabajamos con oro de 18 quilates, diamantes y piedras preciosas naturales de proveedores de confianza. Más de treinta años de relaciones nos avalan.", "Eliminar referencias a certificaciones específicas si no están activas. Enfatizar la trayectoria."),
    ("Nosotros", "Valores", "Valor 3 título", "Transparencia", "Confianza y transparencia", "Mantener el concepto, es un valor real para empresa familiar."),
    ("Nosotros", "Valores", "Valor 3 texto", "Te decimos exactamente qué materiales usamos, de dónde vienen y cuánto valen. El precio que pagas es justo: ni un euro de más, ni uno de menos.", "Como empresa familiar, construimos relaciones de confianza con cada cliente. Te explicamos qué materiales usamos, cómo los trabajamos y cuánto valen. Sin sorpresas.", "Más personal y auténtico para empresa familiar."),
    ("Nosotros", "Valores", "Valor 4 título", "Relación vitalicia", "Una relación para toda la vida", "Mantener el concepto, es coherente con el servicio de reparaciones."),
    ("Nosotros", "Valores", "Valor 4 texto", "Cuando compras en Moreno Santos, no cierras una transacción — comienzas una relación. Mantenemos, reparamos y actualizamos tus joyas para siempre.", "Cuando compras en Moreno Santos Joiers, no cierras una transacción — empiezas una relación. Reparamos, restauramos, grabamos y ponemos al día tus joyas, seas o no hayas sido cliente.", "Añadir los servicios reales: reparar, restaurar, grabar. Abrir al público general no solo clientes."),
    ("Nosotros", "Timeline", "Título", "15 años historia", "Más de 30 años de historia", "Corrección numérica fundamental."),
    ("Nosotros", "Timeline", "Subtítulo", "Cada hito marca el camino de una marca construida con paciencia, pasión y propósito.", "Cada etapa marca el camino de una empresa familiar construida con paciencia, pasión y propósito.", "Sustituir 'marca' por 'empresa familiar'."),
    ("Nosotros", "Timeline", "Hito 1", "Carlos y Lucía abren el primer taller en un local de 40m² en el Barrio Gótico. Solo hacen encargos personales. Los primeros clientes son amigos y familia.", "[Historia real de los fundadores de Moreno Santos Joiers y cómo abrieron el taller en 1993]", "Sustituir por la historia real de la familia fundadora. Los nombres 'Carlos y Lucía' son ficticios."),
    ("Nosotros", "Timeline", "Hito 2", "Lanzan la colección 'Barri' — 12 piezas inspiradas en los colores y texturas del Gótico. Se agota en dos semanas. Aparece en la revista AD España.", "[Primer hito real de la empresa: primer encargo especial, primera colección, reconocimiento importante, etc.]", "Sustituir por hitos reales de la empresa."),
    ("Nosotros", "Timeline", "Hito 3 (Premio)", "La Generalitat de Catalunya les reconoce como uno de los 10 talleres artesanos más relevantes de la comunidad.", "[Premio o reconocimiento real obtenido por la empresa, si existe]", "Solo incluir si es un logro real. Eliminar si es ficticio."),
    ("Nosotros", "Timeline", "Hito 4 (Internacional)", "Primeras piezas que viajan a 20 países. Colaboración con el Hotel Arts Barcelona.", "[Logros reales de expansión o colaboraciones notables]", "Solo incluir datos verificables."),
    ("Nosotros", "Timeline", "Hito 5 (B Corp)", "Obtenemos la certificación de empresa de impacto positivo B Corp.", "[Eliminar si no es real]", "No incluir certificaciones que no se tengan. Puede dañar la credibilidad."),
    ("Nosotros", "Cierre", "Texto final", "850+ piezas únicas creadas, clientes en 38 países y la misma pasión del primer día. Seguimos siendo el mismo taller del Barrio Gótico de siempre.", "Más de 30 años creando joyas únicas, alianzas que unen vidas y recuerdos que duran para siempre. Seguimos siendo la misma joyería familiar de siempre.", "Eliminar datos no verificables. Mensaje auténtico centrado en la familia y la permanencia."),

    # ===================== ATELIER =====================
    ("Atelier", "Hero", "Ubicación", "Barrio Gótico · Barcelona", "Barcelona", "Usar ubicación genérica hasta confirmar dirección real."),
    ("Atelier", "Hero", "Título sección", "Nuestro espacio — donde nace", "Nuestro taller — donde todo nace", "Usar 'taller' en lugar de 'atelier' (más cercano al cliente español/catalán)."),
    ("Atelier", "Hero", "Descripción", "En el corazón del Barrio Gótico, un espacio donde el tiempo se detiene y cada pieza cobra vida entre las manos del artesano.", "En nuestro taller de Barcelona, un espacio donde el tiempo se detiene y cada pieza cobra vida entre las manos del artesano. Más de treinta años de oficio en el mismo lugar.", "Adaptar ubicación y añadir trayectoria temporal."),
    ("Atelier", "Hero", "CTA", "Visitar el atelier", "Visítanos en el taller", "Más cercano y directo."),
    ("Atelier", "Proceso", "Título", "Del boceto a la joya perfecta", "Del boceto a tu joya perfecta", "Personalizar: 'tu joya', más cercano."),
    ("Atelier", "Proceso", "Paso 1 título", "Diseño & Boceto", "Primera conversación y diseño", "Más descriptivo del proceso real."),
    ("Atelier", "Proceso", "Paso 1 texto", "Todo comienza con una conversación. Escuchamos tu visión y la traducimos a trazos únicos, hasta que la forma te emocione.", "Todo empieza con una conversación. Te escuchamos, entendemos qué quieres expresar y lo traducimos a bocetos únicos. Trabajamos contigo hasta que la forma te emocione.", "Más conversacional y próximo."),
    ("Atelier", "Proceso", "Paso 2 título", "Selección de materiales", "Elección de materiales", "Más natural en español."),
    ("Atelier", "Proceso", "Paso 2 texto", "Elegimos cada metal y cada gema personalmente. Solo trabajamos con proveedores de trazabilidad certificada y origen ético.", "Elegimos cada metal y cada gema de forma personal. Trabajamos con oro de 18 quilates, diamantes y piedras preciosas naturales de la máxima calidad.", "Eliminar referencias a certificaciones genéricas. Mencionar materiales reales."),
    ("Atelier", "Proceso", "Paso 3 título", "Elaboración artesanal", "Elaboración en nuestro taller", "Destacar que el trabajo se hace en taller propio."),
    ("Atelier", "Proceso", "Paso 3 texto", "En nuestro atelier, el orfebre trabaja cada pieza a mano, respetando los tiempos que exige la excelencia. Sin atajos.", "En nuestro taller propio, cada pieza se trabaja a mano, respetando los tiempos que exige la excelencia. Sin atajos, sin intermediarios.", "Añadir 'propio' y 'sin intermediarios'."),
    ("Atelier", "Proceso", "Paso 4 título", "Entrega & Garantía", "Entrega y seguimiento", "Más honesto y menos 'garantía vitalicia' si no es política real."),
    ("Atelier", "Proceso", "Paso 4 texto", "Cada joya se entrega con certificado de autenticidad y garantía vitalicia de mantenimiento. Es tuya para siempre.", "Cada joya se entrega con su certificado y toda la información sobre los materiales usados. Y si un día necesitas ajustarla o repararla, aquí seguimos.", "Más real y cercano. Incluir referencia al servicio postventa de reparaciones."),
    ("Atelier", "Diamantes", "Título", "Diamantes certificados — Elige tu diamante", "Diamantes y gemas naturales — Nuestra selección", "Ampliar a todas las gemas que trabajáis."),
    ("Atelier", "Diamantes", "Descripción", "Explora nuestra selección de diamantes certificados. Compara talla, color, pureza y quilataje antes de decidir.", "Trabajamos con diamantes y piedras preciosas naturales seleccionadas. Consúltanos y te ayudamos a elegir la gema perfecta para tu joya.", "Más cercano y sin promesas tecnológicas aún no disponibles."),
    ("Atelier", "Diseñador 3D", "Descripción", "Configura el metal, la talla y la piedra de tu anillo y visualízalo en 3D antes de encargar tu pieza única.", "Cuéntanos qué imaginas y nuestro equipo te preparará una propuesta con bocetos y opciones de materiales antes de empezar.", "Sustituir funcionalidad 3D por el proceso real de consulta y boceto artesanal."),

    # ===================== ENCARGO =====================
    ("Encargo", "Hero", "Título", "Piezas únicas — Tu joya, creada para ti", "Diseño a medida — Tu joya, creada por nosotros", "Más directo: quien crea la joya es vuestro taller."),
    ("Encargo", "Hero", "Descripción", "El servicio de encargo personal de Moreno Santos es la forma más íntima de poseer una joya. Desde la primera idea hasta la entrega final, cada decisión es tuya.", "El servicio de encargo de Moreno Santos Joiers es la forma más especial de tener una joya propia. Desde la primera idea hasta la entrega, trabajamos juntos cada detalle en nuestro taller.", "Añadir nombre completo y referencia al taller propio."),
    ("Encargo", "Proceso", "Paso 1 texto", "Nos reunimos en el atelier o por videollamada. Cuéntanos tu historia, la ocasión, el presupuesto y lo que imaginas. No necesitas saber de joyería — solo necesitas tener algo que quieras expresar.", "Nos reunimos en el taller o por videollamada. Cuéntanos la ocasión, el presupuesto y lo que imaginas. No necesitas saber de joyería — nosotros nos encargamos del resto.", "Simplificar y centrar en la propuesta de valor: vosotros hacéis el trabajo técnico."),
    ("Encargo", "Proceso", "Paso 2 texto", "Nuestro equipo prepara 2-3 bocetos exclusivos para ti, con selección de materiales, gemas y acabados. Los presentamos en persona con muestras físicas.", "Preparamos bocetos personalizados con propuesta de materiales, gemas y acabados. Te los presentamos y ajustamos hasta que la propuesta sea exactamente lo que buscas.", "Más flexible y cercano, sin prometer número exacto de bocetos."),
    ("Encargo", "Proceso", "Paso 3 texto", "Una vez aprobado el diseño, comienza la creación. Te mantenemos informado con fotos del proceso en cada etapa, para que veas cómo nace tu joya.", "Una vez aprobado el diseño, nuestro taller comienza la creación. Te mantenemos informado en cada etapa para que veas cómo nace tu joya.", "Añadir 'nuestro taller'."),
    ("Encargo", "Proceso", "Paso 4 texto", "La entrega se realiza en el atelier, con una presentación especial. Incluye certificado de autenticidad, caja artesanal y garantía vitalicia de mantenimiento. Envío internacional disponible.", "La entrega se hace en el taller, con toda la documentación de los materiales. Una experiencia especial para cerrar el proceso de creación de tu joya.", "Más honesto sobre lo que se incluye."),
    ("Encargo", "Precios", "Tier 1", "500 – 2K | Esencial — Oro 18k o plata 925", "Desde consulta gratuita — Oro 18k, plata y gemas semipreciosas", "Adaptar a la realidad de vuestro negocio. Los precios los defines tú."),
    ("Encargo", "Precios", "Tier 2", "2K – 8K | Signature — Oro 18k o platino", "Joyas con diamantes y piedras preciosas naturales", "Eliminar platino si no es un metal que trabajéis habitualmente."),
    ("Encargo", "Precios", "Tier 3", "Haute Joaillerie — Platino o oro 24k", "[Nivel premium si aplica a vuestra oferta]", "Solo incluir si es real. Evitar términos de alta joyería parisina si no corresponde al posicionamiento."),
    ("Encargo", "Servicios adicionales", "Título", "También te ayudamos con — Servicios de restauración & grabado", "También te ayudamos con — Reparación, restauración y grabado", "Usar los términos reales de los servicios que ofrecéis."),
    ("Encargo", "Servicios adicionales", "Descripción", "Nuestros artesanos no solo crean nuevas joyas — también devuelven la vida a las antiguas y personalizan las tuyas con grabados únicos.", "Nuestro taller no solo crea joyas nuevas — también repara y restaura joyería y relojería, y realizamos grabaciones en oro, plata y otros metales.", "Añadir relojería (reparación de relojes) y los metales en los que hacéis grabaciones."),

    # ===================== COLECCIONES =====================
    ("Colecciones", "Hero", "Subtítulo", "Joyería de autor", "Joyería artesanal · Diseños propios", "Destacar que son diseños propios, diferenciador real."),
    ("Colecciones", "Hero", "Título", "Nuestras Colecciones", "Nuestras Colecciones", "Correcto, mantener."),
    ("Colecciones", "Hero", "Descripción", "Cada línea nace de una emoción distinta. Piezas únicas que combinan la tradición artesana con una visión contemporánea del lujo.", "Cada colección nace de una emoción distinta. Diseños propios en oro de 18 quilates, con diamantes y piedras preciosas naturales, elaborados en nuestro taller de Barcelona.", "Destacar diseños propios, materiales reales y taller propio."),
    ("Colecciones", "Filtros", "Filtro 1", "Collares", "Collares", "Mantener si tenéis esa categoría."),
    ("Colecciones", "Filtros", "Filtro 2", "Pendientes", "Pendientes", "Mantener."),
    ("Colecciones", "Filtros", "Filtro 3", "Pulseras", "Pulseras", "Considerar añadir: Anillos, Alianzas, Plata."),
    ("Colecciones", "Edición limitada", "Título", "Exclusividad — Piezas de edición limitada", "Diseños especiales — Piezas únicas de nuestro taller", "Menos pretencioso, más artesanal."),
    ("Colecciones", "Edición limitada", "Descripción", "Solo 12 unidades de cada diseño. Una vez agotadas, no se repiten. Porque el verdadero lujo es también la rareza.", "Diseños exclusivos creados en nuestro taller. Cuando se acaban, no se repiten — cada pieza tiene su propia historia.", "Mismo concepto pero más cercano y menos 'luxury brand'."),
    ("Colecciones", "CTA encargo", "Texto", "Encargo a medida", "¿No encuentras lo que buscas? Creamos tu joya a medida", "Más útil para el usuario: invita a contactar si no hay nada en la colección que les convenza."),

    # ===================== CONTACTO =====================
    ("Contacto", "Hero", "Título + Subtítulo", "Estamos aquí — Hablemos", "Estamos aquí — Cuéntanos qué necesitas", "Más orientado a la acción y al cliente."),
    ("Contacto", "Hero", "Descripción", "Sea una pregunta sobre una pieza, un encargo personal o simplemente quieras visitarnos — estaremos encantados de atenderte.", "Si tienes una pregunta sobre una joya, quieres encargar algo especial, necesitas una reparación o simplemente quieres visitarnos — estamos encantados de atenderte.", "Añadir 'reparación' como motivo de contacto, servicio real de la empresa."),
    ("Contacto", "Email", "Dirección", "[email protected] (ofuscado Cloudflare)", "morenosantosjoiers@gmail.com", "Email real actualizado en el HTML. Enlace mailto: funcional."),
    ("Contacto", "Teléfono", "Número", "+34 932 000 000", "+34 93 219 00 11", "Teléfono real actualizado en el HTML con href tel:+34932190011."),
    ("Contacto", "WhatsApp", "Número", "+34 600 000 000", "+34 93 219 00 11", "Mismo número que teléfono. Actualizado en el HTML con href wa.me/34932190011."),
    ("Contacto", "Formulario", "Título", "Escríbenos lo que necesitas", "Escríbenos — Te respondemos en menos de 24h", "Añadir tiempo de respuesta en el título para generar confianza."),
    ("Contacto", "Formulario", "Opción dropdown 1", "Información sobre una pieza", "Información sobre una joya", "Usar 'joya' en lugar de 'pieza'."),
    ("Contacto", "Formulario", "Opción dropdown 2", "Encargo personal", "Quiero crear una joya a medida", "Más descriptivo y accionable."),
    ("Contacto", "Formulario", "Opción dropdown 3", "Visita al atelier", "Quiero visitar el taller", "Usar 'taller' en lugar de 'atelier'."),
    ("Contacto", "Formulario", "Opción dropdown 4", "Reparación o mantenimiento", "Reparación de joya o reloj", "Añadir 'reloj' ya que también reparáis relojería."),
    ("Contacto", "Formulario", "Opción dropdown 5", "Colaboración o prensa", "Colaboración o prensa", "Mantener si aplica."),
    ("Contacto", "Formulario", "Opción dropdown 6", "[No existe]", "Grabación personalizada", "Añadir opción para el servicio de grabaciones."),
    ("Contacto", "Ubicación", "Dirección", "Carrer del Bisbe, 12 · Barrio Gótico, Barcelona 08002", "Carrer de l'Escorial, 162, Local 2 · Barcelona 08024", "Dirección real actualizada en el HTML."),
    ("Contacto", "Ubicación", "Metro / barrio", "Barrio Gótico · Metro Jaume I (L4)", "Barcelona 08024 · Local 2", "Actualizado en el mapa del HTML. Añadir indicaciones de metro/bus si procede."),
    ("Contacto", "Horarios", "Lunes-Viernes", "10:00 – 19:00", "[Horario real de apertura]", "Actualizar con horario real."),
    ("Contacto", "Horarios", "Sábado", "11:00 – 18:00", "[Horario real de sábado]", "Actualizar con horario real."),
    ("Contacto", "Servicios", "Reparaciones", "Aceptamos joyas de otras marcas. Consulta disponibilidad y presupuesto sin compromiso.", "Reparamos y restauramos joyas y relojes de cualquier marca. Consulta sin compromiso.", "Añadir 'relojes' explícitamente."),

    # ===================== GLOBAL / TODAS LAS PÁGINAS =====================
    ("Global", "Navegación", "Item 1", "Colecciones", "Colecciones", "Correcto."),
    ("Global", "Navegación", "Item 2", "Atelier", "Taller", "Usar 'Taller' es más natural en español/catalán para empresa familiar."),
    ("Global", "Navegación", "Item 3", "Encargo", "Encargo a medida", "Más descriptivo."),
    ("Global", "Navegación", "Item 4", "Nosotros", "Nosotros", "Correcto."),
    ("Global", "Navegación", "Item 5", "Contacto", "Contacto", "Correcto."),
    ("Global", "Redes sociales", "Instagram handle", "@morenosantos", "[Handle real de Instagram de Moreno Santos Joiers]", "Actualizar con cuenta real."),
    ("Global", "Redes sociales", "Pinterest", "Moreno Santos", "[Cuenta real de Pinterest si existe]", "Verificar si tenéis Pinterest activo."),
    ("Global", "Footer links", "Privacidad", "Privacidad", "Política de privacidad", "Texto completo es más claro para el usuario."),
    ("Global", "Footer links", "Envíos", "Envíos", "Envíos y devoluciones", "Ampliar si también tenéis política de devoluciones."),
    ("Global", "Meta / SEO", "Descripción general", "[No definida en el código]", "Joyería artesanal en Barcelona desde 1993. Diseños propios en oro de 18 quilates, diamantes y piedras preciosas. Taller propio, alianzas de boda, reparaciones y grabaciones.", "Añadir meta description en todas las páginas para mejorar el SEO."),
]

# --- Escribir datos ---
for i, row in enumerate(data, start=2):
    ws.append(row)
    bg = COLOR_ROW_ODD if i % 2 == 0 else COLOR_ROW_EVEN
    for j, cell in enumerate(ws[i], start=1):
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if j == 1:  # Página
            cell.fill = PatternFill("solid", fgColor=COLOR_PAGE_BG)
            cell.font = Font(bold=True, size=10)
        elif j == 4:  # Texto actual
            cell.fill = PatternFill("solid", fgColor=COLOR_CURRENT)
            cell.font = Font(size=10)
        elif j == 5:  # Texto propuesto
            cell.fill = PatternFill("solid", fgColor=COLOR_PROPOSED)
            cell.font = Font(size=10)
        elif j == 6:  # Notas
            cell.fill = PatternFill("solid", fgColor="E8F4F8")
            cell.font = Font(size=9, italic=True, color="555555")
        else:
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.font = Font(size=10)

# --- Anchos de columna ---
col_widths = [14, 22, 28, 55, 55, 50]
for i, width in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = width

# --- Altura de filas ---
for row in ws.iter_rows():
    ws.row_dimensions[row[0].row].height = 15

# --- Congelar primera fila ---
ws.freeze_panes = "A2"

# --- Guardar ---
output_path = "/home/user/MorenoSantosJoiers/textos_web_moreno_santos.xlsx"
wb.save(output_path)
print(f"Excel generado: {output_path}")
print(f"Total filas de texto: {len(data)}")
